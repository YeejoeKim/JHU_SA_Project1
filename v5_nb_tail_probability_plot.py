"""Plot empirical and negative-binomial tail probabilities from the v5 workbook.

The plot compares P(X >= k), so it is a survival (tail-probability) plot,
not a cumulative distribution function (CDF) plot.
"""

from __future__ import annotations

import math
from datetime import time
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


INPUT_FILE = Path(r"E:\Desktop\统计分析\Module 2\Processed_v5_Empirical_Tail_Probabilities.xlsx")
OUTPUT_FILE = Path("outputs/nb_tail_probability_with_parameters.png")

# Four complete years at 10:00–11:00, followed by one high-error comparison.
PLOTS = [
    ("2016", 10),
    ("2017", 10),
    ("2018", 10),
    ("2019", 10),
    ("2018", 13),
]


def nb_tail_probability(k: int, r: float, p: float) -> float:
    """Return P(X >= k) for X ~ NegativeBinomial(r, p).

    Here X counts failures before r successes. This is the same parameterisation
    as r = mu^2/(variance - mu) and p = mu/variance.
    """
    pmf = p**r  # P(X = 0)
    cdf_before_k = pmf
    for x in range(1, k):
        pmf *= ((x + r - 1) / x) * (1 - p)
        cdf_before_k += pmf
    return max(0.0, 1.0 - cdf_before_k)


def hour_value(value: object) -> int | None:
    """Convert an Excel time cell to an hour number."""
    if isinstance(value, time):
        return value.hour
    if isinstance(value, (int, float)):
        return int(value * 24) % 24
    return None


def read_daily_counts(sheet, target_hour: int) -> np.ndarray:
    """Find the daily-count table and return the selected hour's counts."""
    header_row = None
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row, 1).value == "Hour":
            header_row = row
            break
    if header_row is None:
        raise ValueError(f"Could not find the daily-count table in {sheet.title}.")

    hour_column = None
    for column in range(2, sheet.max_column + 1):
        if hour_value(sheet.cell(header_row, column).value) == target_hour:
            hour_column = column
            break
    if hour_column is None:
        raise ValueError(f"Hour {target_hour}:00 was not found in {sheet.title}.")

    counts = []
    for row in range(header_row + 1, sheet.max_row + 1):
        date_value = sheet.cell(row, 1).value
        count = sheet.cell(row, hour_column).value
        if date_value is not None and isinstance(count, (int, float)):
            counts.append(float(count))
    return np.asarray(counts)


def read_saved_actual_probabilities(sheet, target_hour: int, thresholds: np.ndarray) -> np.ndarray:
    """Read the Actual P(X>=k) row already calculated in the Excel file."""
    for row in range(1, 6):
        if hour_value(sheet.cell(row, 1).value) == target_hour:
            values = [sheet.cell(row, int(k) + 1).value for k in thresholds]
            return np.asarray(values, dtype=float)
    raise ValueError(f"Actual-probability row for {target_hour}:00 was not found.")


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    """Use a Windows font if available; otherwise use Pillow's fallback font."""
    path = "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"
    if Path(path).exists():
        return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def draw_panel(draw, bounds, title, thresholds, actual, model) -> None:
    """Draw one tail-probability panel with a step empirical curve."""
    left, top, right, bottom = bounds
    title_font, body_font, small_font = font(18, True), font(16), font(13)
    plot_top = top + 80
    plot_bottom = bottom - 48
    plot_left = left + 58
    plot_right = right - 14

    draw.multiline_text((left, top), title, fill="#243B53", font=title_font, spacing=3)

    def x_pos(k):
        return plot_left + (k - 1) / 49 * (plot_right - plot_left)

    def y_pos(probability):
        return plot_bottom - probability * (plot_bottom - plot_top)

    # Grid and axes
    for value in [0, 0.25, 0.50, 0.75, 1.0]:
        y = y_pos(value)
        draw.line((plot_left, y, plot_right, y), fill="#D9E2EC", width=1)
        draw.text((left + 2, y - 8), f"{value:.0%}", fill="#52606D", font=small_font)
    for tick in [1, 10, 20, 30, 40, 50]:
        x = x_pos(tick)
        draw.line((x, plot_top, x, plot_bottom), fill="#EDF2F7", width=1)
        draw.text((x - 7, plot_bottom + 7), str(tick), fill="#52606D", font=small_font)
    draw.line((plot_left, plot_top, plot_left, plot_bottom), fill="#52606D", width=2)
    draw.line((plot_left, plot_bottom, plot_right, plot_bottom), fill="#52606D", width=2)

    # Actual probability: a post-step curve, because it is empirical/discrete.
    actual_points = []
    for i, k in enumerate(thresholds[:-1]):
        actual_points.extend([(x_pos(k), y_pos(actual[i])), (x_pos(thresholds[i + 1]), y_pos(actual[i]))])
        actual_points.append((x_pos(thresholds[i + 1]), y_pos(actual[i + 1])))
    actual_points.append((x_pos(thresholds[-1]), y_pos(actual[-1])))
    draw.line(actual_points, fill="#17827A", width=4)

    # Negative-binomial tail probability.
    model_points = [(x_pos(k), y_pos(probability)) for k, probability in zip(thresholds, model)]
    draw.line(model_points, fill="#2864B7", width=4)
    draw.text((plot_left + 135, bottom - 27), "Threshold k (calls)", fill="#52606D", font=body_font)


def draw_figure(series) -> None:
    """Create one five-panel, publication-ready PNG without Matplotlib."""
    width, height = 1800, 1160
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 38), "Actual vs negative-binomial tail probabilities for 911 calls",
              fill="#102A43", font=font(32, True))
    draw.line((70, 95, 115, 95), fill="#17827A", width=5)
    draw.text((125, 84), "Actual empirical P(X≥k)", fill="#243B53", font=font(18))
    draw.line((410, 95, 455, 95), fill="#2864B7", width=5)
    draw.text((465, 84), "Negative binomial P(X≥k)", fill="#243B53", font=font(18))

    panels = [
        (70, 150, 600, 570), (640, 150, 1170, 570), (1210, 150, 1740, 570),
        (70, 640, 600, 1060), (640, 640, 1170, 1060),
    ]
    for panel, item in zip(panels, series):
        title = (
            f"{item['year']} | {item['hour']:02d}:00–{item['hour'] + 1:02d}:00 | n={item['n']}\n"
            f"μ={item['mu']:.2f}, s²={item['variance']:.2f}, r={item['r']:.2f}, p={item['p']:.3f}"
        )
        draw_panel(draw, panel, title, item["thresholds"], item["actual"], item["nb_tail"])

    draw.text((70, 1115),
              "Parameters are estimated from each panel's daily counts. Main comparison: 10:00–11:00 in four complete years.",
              fill="#52606D", font=font(16))
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    image.save(OUTPUT_FILE)


def main() -> None:
    if not INPUT_FILE.exists():
        raise FileNotFoundError(f"Cannot find: {INPUT_FILE}")

    # data_only=True reads the calculated Actual P(X>=k) values saved by Excel.
    workbook_values = load_workbook(INPUT_FILE, data_only=True, read_only=True)
    thresholds = np.arange(1, 51)

    series = []
    for year, hour in PLOTS:
        sheet = workbook_values[year]
        counts = read_daily_counts(sheet, hour)
        actual_from_excel = read_saved_actual_probabilities(sheet, hour, thresholds)
        actual_direct = np.array([np.mean(counts >= k) for k in thresholds])

        # A direct calculation from daily counts must equal the saved Excel result.
        difference = float(np.max(np.abs(actual_from_excel - actual_direct)))
        if difference > 1e-12:
            raise ValueError(f"Excel and direct actual probabilities differ by {difference} in {year}, {hour}:00.")

        mu = counts.mean()
        variance = counts.var(ddof=1)
        if variance <= mu:
            raise ValueError(f"Variance must exceed mean for NB fitting: {year}, {hour}:00.")
        r = mu**2 / (variance - mu)
        p = mu / variance
        nb_tail = np.array([nb_tail_probability(int(k), r, p) for k in thresholds])

        series.append({
            "year": year, "hour": hour, "n": len(counts), "mu": mu,
            "variance": variance, "r": r, "p": p,
            "thresholds": thresholds, "actual": actual_direct, "nb_tail": nb_tail,
        })

    draw_figure(series)
    print(f"Saved: {OUTPUT_FILE.resolve()}")
    for item in series:
        print(
            f"{item['year']} {item['hour']:02d}:00 | "
            f"mu={item['mu']:.6f}, s2={item['variance']:.6f}, "
            f"r={item['r']:.6f}, p={item['p']:.6f}"
        )


if __name__ == "__main__":
    main()
